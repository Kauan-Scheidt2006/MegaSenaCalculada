#!/usr/bin/env python3
"""Converte o histórico da Mega-Sena em XLSX para concursos.json."""

from __future__ import annotations

import argparse
import json
import math
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as erro:
    raise SystemExit(
        "Dependência ausente: instale com "
        "'py -m pip install -r requirements.txt'."
    ) from erro


ARQUIVO_PADRAO = Path(
    r"C:\Users\Kauan\Documents\ChatGPT\Mega Sena Calculada\Mega-Sena.xlsx"
)
SAIDA_PADRAO = Path(__file__).resolve().parent / "concursos.json"
NOME_ABA = "MEGA SENA"
CABECALHOS_OBRIGATORIOS = (
    "Concurso",
    "Data do Sorteio",
    "Bola1",
    "Bola2",
    "Bola3",
    "Bola4",
    "Bola5",
    "Bola6",
)


class ErroConversao(ValueError):
    """Indica uma inconsistência nos dados de entrada."""


def inteiro(valor: Any, linha: int, coluna: str) -> int:
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise ErroConversao(
            f"Linha {linha}, coluna '{coluna}': era esperado um inteiro."
        )
    if not math.isfinite(valor) or int(valor) != valor:
        raise ErroConversao(
            f"Linha {linha}, coluna '{coluna}': valor não inteiro ({valor!r})."
        )
    return int(valor)


def formatar_data(valor: Any, linha: int) -> str:
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")

    if isinstance(valor, str):
        texto = valor.strip()
        try:
            return datetime.strptime(texto, "%d/%m/%Y").strftime("%d/%m/%Y")
        except ValueError as erro:
            raise ErroConversao(
                f"Linha {linha}, coluna 'Data do Sorteio': "
                f"data inválida ({valor!r})."
            ) from erro

    raise ErroConversao(
        f"Linha {linha}, coluna 'Data do Sorteio': data ausente ou inválida."
    )


def mapear_cabecalhos(planilha: Any) -> dict[str, int]:
    encontrados: dict[str, int] = {}
    for coluna, celula in enumerate(planilha[1], start=1):
        if isinstance(celula.value, str):
            encontrados[celula.value.strip()] = coluna

    ausentes = [
        cabecalho
        for cabecalho in CABECALHOS_OBRIGATORIOS
        if cabecalho not in encontrados
    ]
    if ausentes:
        raise ErroConversao(
            "Cabeçalhos obrigatórios ausentes: " + ", ".join(ausentes) + "."
        )
    return encontrados


def converter(arquivo: Path, saida: Path) -> int:
    if not arquivo.is_file():
        raise ErroConversao(f"Planilha não encontrada: {arquivo}")

    # O arquivo de origem declara uma dimensão menor que seu conteúdo real.
    # O modo normal do openpyxl recalcula as células existentes corretamente.
    pasta = load_workbook(arquivo, read_only=False, data_only=True)
    try:
        if NOME_ABA not in pasta.sheetnames:
            raise ErroConversao(
                f"Aba '{NOME_ABA}' não encontrada. Abas: {pasta.sheetnames}"
            )

        planilha = pasta[NOME_ABA]
        colunas = mapear_cabecalhos(planilha)
        concursos: list[dict[str, Any]] = []
        numeros_vistos: set[int] = set()

        for numero_linha in range(2, planilha.max_row + 1):
            valores = {
                nome: planilha.cell(numero_linha, coluna).value
                for nome, coluna in colunas.items()
            }

            if all(valor is None for valor in valores.values()):
                continue

            numero_concurso = inteiro(
                valores["Concurso"], numero_linha, "Concurso"
            )
            if numero_concurso <= 0:
                raise ErroConversao(
                    f"Linha {numero_linha}: número do concurso deve ser positivo."
                )
            if numero_concurso in numeros_vistos:
                raise ErroConversao(
                    f"Linha {numero_linha}: concurso {numero_concurso} duplicado."
                )

            dezenas = [
                inteiro(valores[f"Bola{indice}"], numero_linha, f"Bola{indice}")
                for indice in range(1, 7)
            ]
            if any(dezena < 1 or dezena > 60 for dezena in dezenas):
                raise ErroConversao(
                    f"Linha {numero_linha}: dezenas devem estar entre 1 e 60."
                )
            if len(set(dezenas)) != 6:
                raise ErroConversao(
                    f"Linha {numero_linha}: as seis dezenas devem ser distintas."
                )

            concursos.append(
                {
                    "numeroConcurso": numero_concurso,
                    "data": formatar_data(
                        valores["Data do Sorteio"], numero_linha
                    ),
                    "dezenas": dezenas,
                }
            )
            numeros_vistos.add(numero_concurso)
    finally:
        pasta.close()

    concursos.sort(key=lambda concurso: concurso["numeroConcurso"])
    saida.parent.mkdir(parents=True, exist_ok=True)
    with saida.open("w", encoding="utf-8", newline="\n") as arquivo_saida:
        json.dump(concursos, arquivo_saida, ensure_ascii=False, indent=2)
        arquivo_saida.write("\n")

    return len(concursos)


def argumentos() -> argparse.Namespace:
    analisador = argparse.ArgumentParser(
        description="Converte a planilha histórica da Mega-Sena para JSON."
    )
    analisador.add_argument(
        "--entrada",
        type=Path,
        default=ARQUIVO_PADRAO,
        help=f"planilha de origem (padrão: {ARQUIVO_PADRAO})",
    )
    analisador.add_argument(
        "--saida",
        type=Path,
        default=SAIDA_PADRAO,
        help=f"arquivo JSON de destino (padrão: {SAIDA_PADRAO})",
    )
    return analisador.parse_args()


def main() -> int:
    opcoes = argumentos()
    try:
        quantidade = converter(opcoes.entrada, opcoes.saida)
    except (ErroConversao, OSError) as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        return 1

    print(f"{quantidade} concursos gravados em: {opcoes.saida}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
